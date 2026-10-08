"""Importa le schede cotta Excel e il Brewing Schedule nel formato dell'app.

Uso:  python3 tools/importa_storico.py <cartella con gli xlsx> data/storico.json

Le schede cotta hanno tutte lo stesso schema (colonna A: birra, data, lotto, FV,
OG, pH, litri, FG; colonne B..H: sali, acqua mash, acqua sparge, malti, luppoli,
lievito, confezionato; "Note" in fondo). Il Brewing Schedule ha per ogni batch
un blocco di 4 righe: giorno, temperatura, densità e una quarta riga che fino a metà 2025
era la pressione (psi, ~16) e poi è diventata il pH (~4-5.3): sotto 7 è pH, sopra è psi.
"""
import datetime as dt
import glob
import json
import os
import re
import sys

import openpyxl

NUM = r"(\d+(?:[.,]\d+)?)"


def num(s):
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    m = re.search(NUM, str(s))
    return float(m.group(1).replace(",", ".")) if m else None


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def gravity(v):
    """OG/FG scritte come °P (17,1) o come SG x1000 (1062). Ritorna °P."""
    n = num(v)
    if n is None or n == 0:
        return None
    if n > 900:  # 1062 -> SG 1.062
        sg = n / 1000
        return round(-616.868 + 1111.14 * sg - 630.272 * sg**2 + 135.997 * sg**3, 1)
    if 1 < n < 1.2:
        sg = n
        return round(-616.868 + 1111.14 * sg - 630.272 * sg**2 + 135.997 * sg**3, 1)
    return n


UNITS = {"kg": "kg", "gr": "g", "g": "g", "l": "L", "ml": "ml"}


def ingrediente(line):
    """'275KG Pils schuema' / 'magnesio 150gr' / '500gr idaho7 0mn'."""
    s = clean(line)
    if not s:
        return None
    qty, unit, name = None, "", s
    m = re.match(r"^" + NUM + r"\s*(kg|gr|g|ml|l)\b\.?\s*(.*)$", s, re.I)
    if m:
        qty, unit, name = float(m.group(1).replace(",", ".")), UNITS[m.group(2).lower()], m.group(3)
    else:
        m = re.match(r"^(.*?)\s+" + NUM + r"\s*(kg|gr|g|ml|l)\b\.?\s*(.*)$", s, re.I)
        if m:
            name = (m.group(1) + " " + m.group(4)).strip()
            qty, unit = float(m.group(2).replace(",", ".")), UNITS[m.group(3).lower()]
    out = {"nome": name.strip(), "qta": qty, "unita": unit}
    t = re.search(r"(\d+)\s*mn\b", name, re.I)
    if t:
        out["minuti"] = int(t.group(1))
    if re.search(r"\bDH\b|dry|giorno", name, re.I):
        out["uso"] = "dry hop"
    return out


def tacca(line):
    """'885L : 91' / '73° : 149' / '684L : 38/36' -> tacca del serbatoio acqua calda (cm) dopo i ':'"""
    m = re.search(r"[:;]\s*(\d+(?:[.,]\d+)?)", line)
    return float(m.group(1).replace(",", ".")) if m else None


def confezionato(line):
    s = clean(line)
    m = re.match(r"^" + NUM + r"\s*x\s*" + NUM + r"\s*l?\b", s, re.I)
    if not m:
        return None
    n, size = float(m.group(1).replace(",", ".")), float(m.group(2).replace(",", "."))
    if size < 2:
        tipo = "lattina/bottiglia"
    else:
        tipo = "fusto"
    return {"tipo": tipo, "pezzi": int(n), "litri": size}


def scheda(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    a = lambda r: ws.cell(r, 1).value
    col = lambda c: [clean(ws.cell(r, c).value) for r in range(3, ws.max_row + 1)]
    note_row = next((r for r in range(9, ws.max_row + 1) if clean(a(r)).lower().startswith("note")), None)
    last = (note_row or ws.max_row + 1) - 1
    rng = lambda c: [clean(ws.cell(r, c).value) for r in range(3, last + 1) if clean(ws.cell(r, c).value)]

    birra = re.sub(r"^BIRRA\s+", "", clean(a(1)).split(":", 1)[-1].strip(), flags=re.I)
    lotto = clean(a(3)).split(":", 1)[-1].strip().replace(" ", "")
    data = a(2)
    if isinstance(data, dt.datetime):
        data = data.date().isoformat()
    else:
        data = None
    lotto = re.sub(r"^LOTTO", "", lotto, flags=re.I)
    # il numero viene dal nome file (le cotte doppie hanno lotto "1/2/24")
    m = re.search(r"_(\d+)-(\d+)(?:\s*\(\d+\))?\.xlsx$", path)
    numero, anno = (int(m.group(1)), 2000 + int(m.group(2))) if m else (None, None)
    if numero:
        lotto = f"{numero}/{str(anno)[2:]}"  # la cella LOTTO a volte non è aggiornata
    if not data and anno:
        data = f"{anno}-01-01"

    mash, sparge = rng(3), rng(4)
    rec = {
        "id": f"cotta-{lotto.replace('/', '-')}-{birra.lower().replace(' ', '')}",
        "tipo": "cotta",
        "birra": birra,
        "lotto": lotto,
        "numero": numero,
        "anno": anno,
        "data": data,
        "fv": clean(a(4)).upper().replace(" ", ""),
        "og": gravity(clean(a(5))[2:]),
        "fg": gravity(clean(a(8))[2:]),
        "phMash": num(clean(a(6))[2:]),
        "litri": num(a(7)),
        "sali": [x for x in map(ingrediente, rng(2)) if x],
        "malti": [x for x in map(ingrediente, rng(5)) if x],
        "luppoli": [x for x in map(ingrediente, rng(6)) if x],
        "lievito": [x for x in map(ingrediente, rng(7)) if x],
        "acquaMash": {"righe": mash},
        "acquaSparge": {"righe": sparge},
        "confezionato": [x for x in map(confezionato, rng(8)) if x],
        "confezionatoNote": "; ".join(x for x in rng(8) if not confezionato(x)),
        "note": clean(ws.cell(note_row, 2).value) if note_row else "",
        "fermentazione": [],
        "fonte": os.path.basename(path),
    }
    for r in mash:
        if re.match(r"^T°\s*\d", r):
            rec["acquaMash"]["tempMash"] = num(r)
        elif re.search(r"amb|esterna", r, re.I):
            rec["acquaMash"]["tempAmbiente"] = num(r)
        elif re.match(r"^PH", r, re.I):
            rec["acquaMash"]["ph15"] = num(r.split(":")[-1])
        elif re.match(r"^\d+[.,]?\d*\s*L", r):
            rec["acquaMash"]["litri"] = num(r)
            rec["acquaMash"]["taccaFine"] = tacca(r)
        elif re.match(r"^\d+[.,]?\d*°\s*:", r):
            rec["acquaMash"]["tempAcqua"] = num(r)
            rec["acquaMash"]["taccaInizio"] = tacca(r)
    for r in sparge:
        if re.match(r"^\d+[.,]?\d*°$", r):
            rec["acquaSparge"]["temp"] = num(r)
        elif re.match(r"^\d+[.,]?\d*\s*L", r):
            rec["acquaSparge"]["litri"] = num(r)
            rec["acquaSparge"]["taccaFine"] = tacca(r)
    rec["stato"] = "confezionata" if rec["confezionato"] else "chiusa"
    return rec


VERDE_MATERIE_PRIME = "FF92D050"


def schedule(path):
    """Blocchi del Brewing Schedule 2024+ -> {(numero, anno): {...}}"""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        if ws.cell(1, 1).value != "Batch" or not isinstance(ws.cell(1, 5).value, dt.datetime):
            continue
        dates = {c: ws.cell(1, c).value for c in range(5, ws.max_column + 1) if isinstance(ws.cell(1, c).value, dt.datetime)}
        for r in range(2, ws.max_row + 1):
            batch = ws.cell(r, 1).value
            if batch in (None, ""):
                continue
            log, brew = [], None
            for c, d in dates.items():
                day = ws.cell(r, c).value
                if not isinstance(day, (int, float)):
                    continue
                if brew is None:
                    brew = (d - dt.timedelta(days=int(day) - 1)).date()
                e = {"data": d.date().isoformat(), "giorno": int(day)}
                for k, off in (("temp", 1), ("densita", 2), ("psi", 3)):
                    v = ws.cell(r + off, c).value
                    if isinstance(v, (int, float)) and v != 0:
                        e["ph" if k == "psi" and v < 7 else k] = v
                log.append(e)
            if brew is None:
                continue
            fv = clean(ws.cell(r, 4).value)
            fvs = [f"FV{x}" for x in re.findall(r"\d+", fv)]
            for n in re.findall(r"\d+", str(batch)):
                key = (int(n), brew.year)
                prev = out.get(key)
                if prev and len(prev["fermentazione"]) >= len(log):
                    continue
                out[key] = {
                    "birra": re.sub(r"^BIRRA\s+", "", clean(ws.cell(r, 2).value), flags=re.I),
                    "lievitoNome": clean(ws.cell(r, 3).value),
                    "data": brew.isoformat(),
                    "fvPercorso": fvs,
                    "fermentazione": log,
                    "batchRaw": clean(batch),
                    # casella del n° cotta in verde = materie prime ordinate o in magazzino
                    "materiePrime": ws.cell(r, 1).fill.fgColor.rgb == VERDE_MATERIE_PRIME,
                }
    return out


def chiudi_fermentatori(cotte):
    """Data di fine in FV: ultimo giorno con temperatura impostata nello schedule.
    Lo schedule continua spesso a contare i giorni anche dopo lo svuotamento, quindi
    per le cotte passate la fine si ferma al giorno prima della cotta successiva nello
    stesso FV (le cotte doppie della stessa birra contano come una sola)."""
    oggi = dt.date.today().isoformat()
    per_fv = {}
    for c in cotte:
        if c.get("fv") and c.get("data"):
            per_fv.setdefault(c["fv"], []).append(c)
    for lista in per_fv.values():
        lista.sort(key=lambda c: c["data"])
        for i, c in enumerate(lista):
            log = [e for e in c["fermentazione"] if e.get("temp")]
            fine = log[-1]["data"] if log else None
            nxt = next((n for n in lista[i + 1:] if n["data"] > c["data"] and not (
                n["birra"].upper() == c["birra"].upper()
                and (dt.date.fromisoformat(n["data"]) - dt.date.fromisoformat(c["data"])).days <= 2)), None)
            if nxt and nxt["data"] <= oggi:
                limite = (dt.date.fromisoformat(nxt["data"]) - dt.timedelta(days=1)).isoformat()
                if not fine or fine > limite:
                    fine = limite
            if fine and fine < c["data"]:
                fine = c["data"]
            if fine:
                c["fine"] = fine
                c["fermentazione"] = [e for e in c["fermentazione"] if e["data"] <= fine and len(e) > 2]


def main(src, dst):
    cotte = []
    for f in sorted(glob.glob(os.path.join(src, "SCHEDA-COTTA*.xlsx"))):
        try:
            cotte.append(scheda(f))
        except Exception as e:  # noqa: BLE001
            print("ERRORE", f, e, file=sys.stderr)
    sched_path = os.path.join(src, "Brewing Schedule Kashmir.xlsx")
    sched = schedule(sched_path) if os.path.exists(sched_path) else {}
    used = set()
    for c in cotte:
        s = sched.get((c["numero"], c["anno"]))
        if s:
            c["fermentazione"] = s["fermentazione"]
            c["data"] = s["data"]  # la data della scheda a volte è quella del modello copiato
            if len(s["fvPercorso"]) > 1:
                c["fvPercorso"] = s["fvPercorso"]
            if s["materiePrime"]:
                c["materiePrime"] = True
            used.add((c["numero"], c["anno"]))
    # batch presenti solo nello schedule (senza scheda cotta)
    for (n, anno), s in sorted(sched.items()):
        if (n, anno) in used:
            continue
        cotte.append({
            "id": f"cotta-{n}-{str(anno)[2:]}-sched",
            "tipo": "cotta", "birra": s["birra"].strip(), "lotto": f"{n}/{str(anno)[2:]}",
            "numero": n, "anno": anno, "data": s["data"],
            "fv": s["fvPercorso"][0] if s["fvPercorso"] else "",
            "fvPercorso": s["fvPercorso"] if len(s["fvPercorso"]) > 1 else [],
            "lievito": [{"nome": s["lievitoNome"], "qta": None, "unita": ""}] if s["lievitoNome"] else [],
            "og": None, "fg": None, "phMash": None, "litri": None,
            "sali": [], "malti": [], "luppoli": [], "confezionato": [], "confezionatoNote": "",
            "acquaMash": {"righe": []}, "acquaSparge": {"righe": []},
            "fermentazione": s["fermentazione"], "note": "",
            **({"materiePrime": True} if s["materiePrime"] else {}),
            "stato": "pianificata" if s["data"] >= dt.date.today().isoformat() else "chiusa",
            "fonte": "Brewing Schedule Kashmir.xlsx",
        })
    chiudi_fermentatori(cotte)
    seen = {}
    for c in cotte:
        while c["id"] in seen:
            c["id"] += "x"
        seen[c["id"]] = 1
    cotte.sort(key=lambda c: (c["data"] or "", c["numero"] or 0))
    with open(dst, "w") as fh:
        json.dump({"versione": 1, "generato": dt.date.today().isoformat(), "cotte": cotte}, fh, ensure_ascii=False)
    print(f"{len(cotte)} cotte ({len(used)} con log fermentazione dallo schedule, {len(sched) - len(used)} solo da schedule)")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
